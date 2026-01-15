from __future__ import annotations

import io
import json
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Dict

from docx import Document
from flask import Flask, jsonify, request, send_file, send_from_directory

app = Flask(__name__, static_folder="static", static_url_path="/static")

DATE_PATTERN = re.compile(r"^(\d{4}[._-]\d{1,2}[._-]\d{1,2})\s*")
DEFAULT_REPLACEMENTS: Dict[str, str] = {
    "2024.12.01": "2026.1.30",
    "示例签名": "张三",
    "示例公司": "XX 公司",
}


@app.route("/")
def home() -> "flask.wrappers.Response":
    return send_from_directory(Path(__file__).parent, "index.html")


@app.route("/api/rename", methods=["POST"])
def rename_files():
    upload = request.files.get("file")
    new_date = request.form.get("date", "2026.1.30")
    separator = request.form.get("separator", " ")

    if not upload:
        return jsonify({"error": "缺少文件，先上传包含 .docx 的压缩包"}), 400

    with tempfile.TemporaryDirectory() as tmpdir:
        zip_path = Path(tmpdir) / "input.zip"
        upload.save(zip_path)

        extract_root = Path(tmpdir) / "input"
        extract_root.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(zip_path, "r") as zf:
            zf.extractall(extract_root)

        renamed_count = 0
        for path in extract_root.rglob("*.docx"):
            original_name = path.name
            stripped = DATE_PATTERN.sub("", original_name)
            new_name = f"{new_date}{separator}{stripped}" if stripped else f"{new_date}{separator}{original_name}"
            target = path.with_name(new_name)
            if target != path:
                path.rename(target)
                renamed_count += 1

        output_stream = io.BytesIO()
        with zipfile.ZipFile(output_stream, "w", zipfile.ZIP_DEFLATED) as zf:
            for file_path in extract_root.rglob("*"):
                if file_path.is_file():
                    zf.write(file_path, file_path.relative_to(extract_root))
        output_stream.seek(0)

    headers = {"X-Renamed": str(renamed_count)}
    return send_file(
        output_stream,
        mimetype="application/zip",
        as_attachment=True,
        download_name="renamed_docs.zip",
        headers=headers,
    )


@app.route("/api/replace", methods=["POST"])
def replace_content():
    upload = request.files.get("file")
    replacements_raw = request.form.get("replacements", "")
    suffix = request.form.get("suffix", "_updated")

    if not upload:
        return jsonify({"error": "缺少文件，先上传包含 .docx 的压缩包"}), 400

    replacements = DEFAULT_REPLACEMENTS
    if replacements_raw:
        try:
            replacements = json.loads(replacements_raw)
        except json.JSONDecodeError:
            return jsonify({"error": "替换映射需为有效 JSON，例如 {\"旧文本\": \"新文本\"}"}), 400

    with tempfile.TemporaryDirectory() as tmpdir:
        zip_path = Path(tmpdir) / "input.zip"
        upload.save(zip_path)

        extract_root = Path(tmpdir) / "input"
        extract_root.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(zip_path, "r") as zf:
            zf.extractall(extract_root)

        processed = 0
        for path in extract_root.rglob("*.docx"):
            doc = Document(path)
            _replace_in_paragraphs(doc.paragraphs, replacements)
            _replace_in_tables(doc.tables, replacements)
            output_path = path.with_name(path.stem + suffix + path.suffix)
            doc.save(output_path)
            processed += 1

        output_stream = io.BytesIO()
        with zipfile.ZipFile(output_stream, "w", zipfile.ZIP_DEFLATED) as zf:
            for file_path in extract_root.rglob("*"):
                if file_path.is_file():
                    zf.write(file_path, file_path.relative_to(extract_root))
        output_stream.seek(0)

    headers = {"X-Processed": str(processed)}
    return send_file(
        output_stream,
        mimetype="application/zip",
        as_attachment=True,
        download_name="replaced_docs.zip",
        headers=headers,
    )


def _replace_in_paragraphs(paragraphs, mapping: Dict[str, str]) -> None:
    for paragraph in paragraphs:
        for old, new in mapping.items():
            if old in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(old, new)


def _replace_in_tables(tables, mapping: Dict[str, str]) -> None:
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                _replace_in_paragraphs(cell.paragraphs, mapping)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
